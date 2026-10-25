import csv,io,secrets
from decimal import Decimal
from openpyxl import load_workbook
import pymupdf
from test_app import mongo_backend,workspace,payload
from app import SALES_REPORT_TYPES
from accounting_reports import REPORT_TYPES
from report_totals import total_row,report_period


def test_report_totals_exact_and_meaningful():
    assert total_row('Sales',['Invoice','Date','Customer','Total','Received','Balance'],[['A','','','89999999999100.01','0','89999999999100.01'],['B','','','0.02','0','0.02']])[3]=='89999999999100.03'
    assert total_row('GST summary',['Source','Taxable','CGST','SGST','IGST'],[['Sales','100','9','9','0'],['Purchases','50','4.50','4.50','0']])[1:]==['50.00','4.50','4.50','0.00']
    assert total_row('Inventory',['Product','SKU','Category','Stock','Minimum','Unit'],[['A','','','2','1','PCS'],['B','','','3','1','BOX']])[3]=='3 BOX; 2 PCS'
    assert report_period({'from':'','to':''},'2026-10-25','Daily sales')==('2026-10-25','2026-10-25')


def test_every_report_runs_with_blank_dates_and_totals_in_all_exports(workspace):
    from test_trade import supplier,purchase
    client,post,backend=workspace
    post('/billing',payload()|dict(received='50'))
    supplier(post);post('/purchases',purchase())
    post('/sales-returns',dict(invoice='1',date='2026-10-18',quantity_0='1',reason='Service exchange',submission=secrets.token_hex(32)))
    post('/purchases/1',dict(action='return',date='2026-10-18',quantity_0='1',reason='Damaged',submission=secrets.token_hex(32)))
    assert backend.database.sales_returns.count_documents({})==1
    assert backend.database.trade_returns.count_documents({})==1
    for kind in SALES_REPORT_TYPES+REPORT_TYPES:
        expected='Net sales' if kind=='GST summary' else 'Total ('
        for export in ('','csv','pdf','xlsx'):
            response=client.get('/reports',query_string={'report':kind,'from':'','to':'','export':export},follow_redirects=True)
            assert response.status_code==200,(kind,export,response.data[:100])
            if export=='pdf':
                with pymupdf.open(stream=response.data,filetype='pdf') as doc: assert expected in ''.join(p.get_text() for p in doc),(kind,export)
            elif export=='xlsx':
                sheet=load_workbook(io.BytesIO(response.data)).active
                assert str(sheet.cell(sheet.max_row,1).value).startswith(expected),(kind,export)
            elif export=='csv':
                rows=list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
                assert rows[-1][0].startswith(expected),(kind,export)
            else: assert expected.encode() in response.data,(kind,export)
    response=client.get('/reports?report=Sales&export=csv&from=2026-10-01&to=2026-10-31')
    rows=list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
    assert rows[-1][3:]==['236.00','50.00','68.00']
    response=client.get('/reports?report=Purchases&export=csv&from=2026-10-01&to=2026-10-31',follow_redirects=True)
    rows=list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
    assert rows[-1][3:]==['180.00','16.20','16.20','0.00','212.40']


def test_report_date_validation_and_empty_totals(workspace):
    client,post,backend=workspace
    for kind in ('Sales','Purchases'):
        for dates in ({'from':'invalid'},{'from':'2026-11-01','to':'2026-01-01'}):
            response=client.get('/reports',query_string=dict(report=kind,**dates),follow_redirects=True)
            assert response.status_code==400
        response=client.get('/reports',query_string=dict(report=kind,export='csv',**{'from':'2000-01-01','to':'2000-12-31'}),follow_redirects=True)
        assert 'Total (0 records)' in response.data.decode('utf-8-sig')
